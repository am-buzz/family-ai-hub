"""Tests for office_export.py using stand-ins for Open WebUI's file storage.
Run from the repo root:  python3 tools/test_office_export.py [output_dir]
Needs: pip install python-docx openpyxl python-pptx
If an output_dir is given, the generated files are kept there so they can be opened and inspected.
"""
import asyncio, io, os, sys, tempfile, types, zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
OUT.mkdir(parents=True, exist_ok=True)
saved = {}

# --- stand-ins for the Open WebUI modules the tool imports ---
files_mod, storage_mod = types.ModuleType("open_webui.models.files"), types.ModuleType("open_webui.storage.provider")
class FileForm:
    def __init__(self, **kw): self.__dict__.update(kw)
class _Files:
    @staticmethod
    async def insert_new_file(user_id, form):  # the real one is async in 0.12
        saved[form.id] = {"user": user_id, "form": form}
        return form
class _Storage:
    @staticmethod
    def upload_file(fh, filename, tags):
        data = fh.read()
        p = OUT / filename
        p.write_bytes(data)
        return data, str(p)
files_mod.FileForm, files_mod.Files, storage_mod.Storage = FileForm, _Files, _Storage
for n, m in [("open_webui", types.ModuleType("open_webui")), ("open_webui.models", types.ModuleType("open_webui.models")),
             ("open_webui.storage", types.ModuleType("open_webui.storage")), ("open_webui.models.files", files_mod), ("open_webui.storage.provider", storage_mod)]:
    sys.modules[n] = m

import office_export as oe
tools = oe.Tools()
events = []
async def emit(e): events.append(e)
user = {"id": "u1", "name": "Anne-Marie", "email": "am@example.com"}
ok = 0
def check(name, cond):
    global ok
    print(("PASS " if cond else "FAIL ") + name); ok += bool(cond)
    if not cond: sys.exit(1)

DOC = """# Briefing: The Swiss Cloud Debate

## Summary
- Switzerland is debating whether public bodies should buy **cloud services** from US providers or from Swiss companies.
- The federal IT tender is worth about *CHF 7 billion*.
  - Sub-point one with a [link](https://www.admin.ch).
  - Sub-point two.
- A decision is expected in 2027.

> **Key takeaway:** Sovereignty and cost pull in opposite directions, and no option is free of trade-offs.

## Options compared
| Option | Provider type | Cost (CHF m) | Data location |
|---|---|---:|:---:|
| A | US hyperscaler | 1,200 | EU/US |
| B | Swiss provider | 1,650 | Switzerland |
| C | Hybrid | 1,400 | Mixed |
| Total | | 4,250 | |

### Risks
1. Vendor lock-in.
2. Legal exposure under foreign law.
3. Capacity of Swiss providers.

> Warning: figures are early estimates and may change.

---

## Sources
- [Federal Chancellery](https://www.bk.admin.ch)
- `inline code` works too
```
code block line
```
"""
r = asyncio.run(tools.create_word_document("Briefing: The Swiss Cloud Debate", DOC, "Prepared for the family", __user__=user, __event_emitter__=emit))
check("word: reports success with a download link", "Done." in r and "/api/v1/files/" in r and ".docx" in r)
fid = list(saved)[-1]
check("word: file registered for the right user with the right type", saved[fid]["user"] == "u1" and saved[fid]["form"].meta["content_type"] == oe.DOCX_MIME)
from docx import Document
d = Document(saved[fid]["form"].path)
styles = [p.style.name for p in d.paragraphs]
check("word: title, headings, list styles present", "Title" in styles and styles.count("Heading 2") >= 3 and "List Bullet" in styles and "List Bullet 2" in styles)
check("word: tables (data table + two call-outs)", len(d.tables) == 3)
data_t = [t for t in d.tables if t.rows[0].cells[0].text == "Option"][0]
check("word: data table has its header and 5 rows", len(data_t.rows) == 5 and data_t.rows[4].cells[0].text == "Total")
check("word: first H1 equal to title is not repeated", sum(1 for p in d.paragraphs if p.text == "Briefing: The Swiss Cloud Debate") == 1)
check("word: hyperlink relationships created", any("admin.ch" in rel.target_ref for rel in d.part.rels.values() if rel.is_external))
check("word: status event emitted", any("Created" in e["data"]["description"] for e in events))

XL = """# Household Budget 2026

## Monthly
| Item | Amount | Share | Due |
|---|---|---|---|
| Rent | CHF 2,400 | 48% | 2026-11-01 |
| Food | $1,200.50 | 24% | 2026-11-05 |
| Phone | 85 | 2% | 2026-11-10 |
| Zip | 01234 | | |
| Total | =SUM(B2:B4) | 74% | |

## Notes sheet
Some text above the table.

| Name | Score |
|---|---|
| Ana | 9 |
| Ben | 7.5 |
| Evil | =HYPERLINK("http://x.test","click") |
"""
r = asyncio.run(tools.create_excel_workbook("Household Budget 2026", XL, __user__=user, __event_emitter__=emit))
check("excel: success + link", "Done." in r and ".xlsx" in r)
fid = list(saved)[-1]
import openpyxl
wb = openpyxl.load_workbook(saved[fid]["form"].path)
check("excel: two sheets with right names", wb.sheetnames == ["Monthly", "Notes sheet"])
ws = wb["Monthly"]
rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
check("excel: percent, date and zip handled", ws.cell(rows["Rent"], 3).value == 0.48 and ws.cell(rows["Rent"], 4).value.year == 2026 and ws.cell(rows["Zip"], 2).value == "01234")
check("excel: currency becomes a number with a format", ws.cell(rows["Rent"], 2).value == 2400 and "CHF" in ws.cell(rows["Rent"], 2).number_format)
check("excel: safe formula kept and takes the column currency format", ws.cell(rows["Total"], 2).value == "=SUM(B2:B4)" and "#,##0.00" in ws.cell(rows["Total"], 2).number_format)
check("excel: freeze panes and filter set", ws.freeze_panes is not None and ws.auto_filter.ref is not None)
ws2 = wb["Notes sheet"]
vals = [c.value for row in ws2.iter_rows() for c in row if c.value]
evil = [c for row in ws2.iter_rows() for c in row if str(c.value or "").startswith("=HYPERLINK")]
check("excel: dangerous formula stored as plain text, not a formula", len(evil) == 1 and evil[0].data_type == "s")
check("excel: table starts at row 1 so B2:B4 means rows 2 to 4", ws.cell(1, 1).value == "Item" and ws.cell(2, 1).value == "Rent")
check("excel: notes written above a table end up below it", ws2.cell(1, 1).value == "Name" and any(c.value == "Some text above the table." for row in ws2.iter_rows() for c in row))

PP = """# Section One

## Why it matters
- Data sovereignty
  - Who can compel access?
- Cost and capacity
- Speed of delivery
Notes: Start with the sovereignty point, then ask the audience.

## The big picture
> Sovereignty and cost pull in opposite directions.

## Comparison
| Option | Cost | Data |
|---|---:|---|
| A | 1,200 | US |
| B | 1,650 | CH |

## A long slide
""" + "\n".join(f"- Point number {i} with a fairly long sentence so that the slide must split into several parts safely" for i in range(1, 12))
r = asyncio.run(tools.create_powerpoint("The Swiss Cloud Debate", PP, "A short briefing", __user__=user, __event_emitter__=emit))
check("pptx: success + link", "Done." in r and ".pptx" in r)
fid = list(saved)[-1]
from pptx import Presentation
pr = Presentation(saved[fid]["form"].path)
check("pptx: title + section + 4 topic slides (the long one splits)", len(pr.slides) >= 8)
titles = [s.shapes.title.text_frame.text for s in pr.slides if s.shapes.title is not None]
check("pptx: titles present and continuation labelled", "Why it matters" in titles and any("(cont.)" in t for t in titles))
check("pptx: speaker notes kept", any(s.has_notes_slide and "sovereignty point" in s.notes_slide.notes_text_frame.text for s in pr.slides))
check("pptx: a table slide exists", any(sh.has_table for s in pr.slides for sh in s.shapes))

# errors are returned to the model as text, never raised
long = "x" * (oe.Tools().valves.max_input_chars + 1)
r = asyncio.run(tools.create_word_document("T", long, "", __user__=user))
check("too-long input returns a readable error", r.startswith("Could not create the Word document"))
check("filenames are safe", oe.slug('Bad/Name: "x"?*') == "BadName x")
# 10 neutral by default; labels and flair are opt-in; Excel keeps its own look
check("word: no label or byline by default", "BUZATU" not in " ".join(p.text.upper() for p in d.paragraphs) and "Prepared by" not in " ".join(p.text for p in d.paragraphs))
check("word: headings use clean Calibri type", d.styles["Heading 1"].font.name == "Calibri Light" and d.styles["Title"].font.name == "Calibri Light")
check("word: headings are blue-gray", str(d.styles["Heading 1"].font.color.rgb) == "3E5C76")
r = asyncio.run(tools.create_word_document("Labelled", "## A\ntext", label="Acme Team", __user__=user))
d2 = Document(list(saved.values())[-1]["form"].path)
check("word: an optional per-document label appears when asked for", any(p.text == "ACME TEAM" for p in d2.paragraphs))
ws = openpyxl.load_workbook([v for v in saved.values() if v["form"].filename.endswith(".xlsx")][0]["form"].path)["Monthly"]
check("excel keeps its original header colour", ws.cell(1, 1).fill.fgColor.rgb.endswith("2A6A98"))
check("excel has no label in the footer unless asked", not ws.oddFooter.right.text)
print(f"\nAll {ok} checks passed. Files are in {OUT}")
