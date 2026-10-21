"""
title: Office Export (Word, Excel, PowerPoint)
author: family-ai-hub
version: 0.1.0
description: Lets a model turn its answer into a polished Word document, Excel workbook or PowerPoint deck and hand back a download link. Files are built on this Mac with python-docx, openpyxl and python-pptx; nothing is sent anywhere.
"""

import asyncio
import datetime as dt
import inspect
import io
import re
import uuid
from typing import Literal, Optional

from pydantic import BaseModel, Field

# =====================================================================================================
# Look and feel. "Lake & Ridge": lake blue, alpine red, ridge green, on warm white. Colours are hex strings.
# =====================================================================================================
INK = "1F2A37"
MUTED = "6B7785"
RULE = "C9D6E2"
TINT = "EAF1F7"
ZEBRA = "F4F8FB"
GREEN = "587F5A"
WARM_FILL = "FDF1E7"
WARM_EDGE = "E07B1F"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# =====================================================================================================
# Markdown reading (a deliberately small subset that models write reliably)
# =====================================================================================================
INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\w*])\*(?![\s*]).+?(?<![\s*])\*(?![\w*])|(?<!\w)_(?![\s_]).+?(?<![\s_])_(?!\w)|`[^`]+`|\[[^\]]+\]\([^)\s]+\))"
)
SAFE_URL = re.compile(r"^(https?://|mailto:)", re.I)


def inline_runs(text: str):
    """Splits text into (text, {'b','i','code','url'}) pieces."""
    text = (text or "").replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    text = re.sub(r"\\([*_`\[\]()#|])", r"\1", text)
    out, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos : m.start()], {}))
        tok = m.group(0)
        if tok.startswith("***"):
            out.append((tok[3:-3], {"b": True, "i": True}))
        elif tok.startswith("**"):
            out.append((tok[2:-2], {"b": True}))
        elif tok.startswith("`"):
            out.append((tok[1:-1], {"code": True}))
        elif tok.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", tok).groups()
            out.append((label, {"url": url} if SAFE_URL.match(url) else {}))
        else:
            out.append((tok[1:-1], {"i": True}))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], {}))
    return [(t, f) for t, f in out if t]


def plain(text: str) -> str:
    return "".join(t for t, _ in inline_runs(text))


def split_row(line: str):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells, cur, esc = [], "", False
    for ch in line:
        if esc:
            cur += ch if ch == "|" else "\\" + ch
            esc = False
        elif ch == "\\":
            esc = True
        elif ch == "|":
            cells.append(cur.strip())
            cur = ""
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


SEP_CELL = re.compile(r"^:?-{2,}:?$")


def parse_blocks(md: str):
    """Returns blocks: ('heading', level, text) ('para', text) ('li', level, ordered, text) ('table', rows, aligns)
    ('quote', text) ('code', text) ('rule',)."""
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks, i = [], 0
    li_re = re.compile(r"^(\s*)([-*+•])\s+(.*)$")
    ol_re = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")

    def starts_block(s: str, nxt: Optional[str]) -> bool:
        st = s.strip()
        return bool(
            not st
            or st.startswith(("```", ">", "|"))
            or re.match(r"^#{1,6}\s", st)
            or re.match(r"^(-{3,}|\*{3,}|_{3,})$", st)
            or li_re.match(s)
            or ol_re.match(s)
        )

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            j, buf = i + 1, []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                buf.append(lines[j])
                j += 1
            blocks.append(("code", "\n".join(buf)))
            i = j + 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*$", s)
        if m:
            blocks.append(("heading", len(m.group(1)), m.group(2)))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            blocks.append(("rule",))
            i += 1
            continue
        if s.startswith("|") or ("|" in s and i + 1 < len(lines) and all(SEP_CELL.match(c) for c in split_row(lines[i + 1]) if c)
                                 and "-" in lines[i + 1] and "|" in lines[i + 1]):
            rows, aligns = [], []
            while i < len(lines) and "|" in lines[i]:
                cells = split_row(lines[i])
                if cells and all(SEP_CELL.match(c) for c in cells if c) and any(cells):
                    aligns = [("c" if c.startswith(":") and c.endswith(":") else "r" if c.endswith(":") else "l") for c in cells]
                else:
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows, aligns))
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append(("quote", "\n".join(buf).strip()))
            continue
        m = li_re.match(raw) or ol_re.match(raw)
        if m:
            ordered = bool(ol_re.match(raw)) and not li_re.match(raw)
            indent = len(m.group(1).replace("\t", "    "))
            level = min(indent // 2, 2)
            text = m.group(3)
            i += 1
            while i < len(lines) and lines[i].strip() and not starts_block(lines[i], None) and lines[i].startswith(" "):
                text += " " + lines[i].strip()
                i += 1
            blocks.append(("li", level, ordered, text))
            continue
        buf = [s]
        i += 1
        while i < len(lines) and not starts_block(lines[i], None):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("para", " ".join(buf)))
    return blocks


def slug(text: str, fallback: str = "document") -> str:
    s = re.sub(r"[^\w\s.-]", "", plain(text), flags=re.U).strip()
    s = re.sub(r"\s+", " ", s)[:80].strip(" .")
    return s or fallback


NUMERIC = re.compile(r"^[\s$€£¥]*[-+]?[\d.,]+\s*%?\s*(CHF|USD|EUR)?\s*$", re.I)


def clamp(n, lo, hi):
    return max(lo, min(hi, n))


# =====================================================================================================
# Word
# =====================================================================================================
def build_docx(title, subtitle, content, author, valves) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor, Twips

    accent, hi = valves.accent_color, valves.highlight_color
    rgb = lambda h: RGBColor.from_string(h.upper())

    doc = Document()
    sec = doc.sections[0]
    if valves.paper_size == "Letter":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    # ---- helpers -------------------------------------------------------------------------------
    def set_font(style, name, size, color=None, bold=None, italic=None):
        style.font.name = name
        style.font.size = Pt(size)
        if color:
            style.font.color.rgb = rgb(color)
        if bold is not None:
            style.font.bold = bold
        if italic is not None:
            style.font.italic = italic
        rpr = style.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            if rf.get(qn("w:" + a)) is not None:
                del rf.attrib[qn("w:" + a)]
        for a in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + a), name)

    def strip_borders(style):
        ppr = style.element.pPr
        if ppr is not None:
            for el in ppr.findall(qn("w:pBdr")):
                ppr.remove(el)

    def para_border(paragraph, side="bottom", color=RULE, size=6, space=2):
        ppr = paragraph._p.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            ppr.append(bdr)
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        bdr.append(el)

    def shade(cell_or_par_pr, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell_or_par_pr.append(shd)

    def cell_shade(cell, fill):
        shade(cell._tc.get_or_add_tcPr(), fill)

    def cell_borders(cell, **edges):
        tcpr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            spec = edges.get(side)
            el = OxmlElement(f"w:{side}")
            if spec:
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(spec[0]))
                el.set(qn("w:color"), spec[1])
            else:
                el.set(qn("w:val"), "nil")
            b.append(el)
        tcpr.append(b)

    def cell_margins(table, top=70, bottom=70, left=110, right=110):
        tblpr = table._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(v))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblpr.append(mar)

    def table_borders(table, color=RULE):
        tblpr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            if side in ("top", "bottom", "insideH"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), color)
            else:
                el.set(qn("w:val"), "nil")
            b.append(el)
        tblpr.append(b)

    def row_flags(row, header=False):
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trpr.append(cs)
        if header:
            th = OxmlElement("w:tblHeader")
            trpr.append(th)

    def field(paragraph, code, size, color):
        """A page-number style field built from separate runs so every part gets the same font."""
        def mk():
            r = paragraph.add_run()
            r.font.size, r.font.color.rgb = Pt(size), rgb(color)
            return r
        for kind in ("begin", "instr", "separate", "text", "end"):
            r = mk()
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {code} "
                r._r.append(el)
            elif kind == "text":
                r.text = "1"
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                r._r.append(el)

    def hyperlink(paragraph, url, text, size=None):
        r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), accent)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(c)
        rpr.append(u)
        if size:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(size * 2)))
            rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        paragraph._p.append(h)

    def add_runs(paragraph, text, size=None, color=None, bold=None, italic=None):
        for t, f in inline_runs(text):
            if f.get("url"):
                hyperlink(paragraph, f["url"], t, size)
                continue
            r = paragraph.add_run(t)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
            if bold or f.get("b"):
                r.font.bold = True
            if italic or f.get("i"):
                r.font.italic = True
            if f.get("code"):
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt((size or 11) - 1)
                r.font.color.rgb = rgb("9A2B20")
                shade(r._r.get_or_add_rPr(), "F1F3F5")
        return paragraph

    # ---- styles --------------------------------------------------------------------------------
    st = doc.styles
    set_font(st["Normal"], "Calibri", 11, INK)
    st["Normal"].paragraph_format.space_after = Pt(7)
    st["Normal"].paragraph_format.line_spacing = 1.18
    for name, font, size, color, before, after in (
        ("Heading 1", "Georgia", 20, accent, 22, 8),
        ("Heading 2", "Georgia", 15, accent, 16, 5),
        ("Heading 3", "Calibri", 12.5, INK, 12, 3),
    ):
        set_font(st[name], font, size, color, bold=True, italic=False)
        pf = st[name].paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        strip_borders(st[name])
    set_font(st["Title"], "Georgia", 28, INK, bold=True)
    strip_borders(st["Title"])
    st["Title"].paragraph_format.space_after = Pt(4)
    st["Title"].paragraph_format.line_spacing = 1.0
    set_font(st["Subtitle"], "Calibri", 14.5, MUTED, italic=False, bold=False)
    st["Subtitle"].paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        set_font(st[name], "Calibri", 11, INK)
        st[name].paragraph_format.space_after = Pt(3)

    # ---- title block ---------------------------------------------------------------------------
    author_line = " · ".join(x for x in (f"Prepared by {author}" if author else "", f"{dt.date.today().day} {dt.date.today():%B %Y}") if x)
    if title:
        kick = doc.add_paragraph()
        kick.paragraph_format.space_after = Pt(6)
        kr = kick.add_run((valves.header_label or "").upper())
        kr.font.size, kr.font.bold = Pt(9), True
        kr.font.color.rgb = rgb(hi)
        kr._r.get_or_add_rPr().append(_spacing(OxmlElement, qn, 50))
        tp = doc.add_paragraph(plain(title), style="Title")
        if subtitle:
            doc.add_paragraph(plain(subtitle), style="Subtitle")
        meta = doc.add_paragraph()
        meta.paragraph_format.space_after = Pt(14)
        mr = meta.add_run(author_line)
        mr.font.size, mr.font.color.rgb = Pt(9.5), rgb(MUTED)
        para_border(meta, "bottom", accent, size=18, space=8)

    # ---- header / footer -----------------------------------------------------------------------
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(plain(title) if title else (valves.header_label or ""))
    hr.font.size, hr.font.color.rgb = Pt(8.5), rgb(MUTED)
    for footer in (sec.footer, sec.first_page_footer):
        fp = footer.paragraphs[0]
        for old in (Twips(4320), Twips(4680), Twips(8640), Twips(9360)):  # the template's own footer tab stops
            fp.paragraph_format.tab_stops.add_tab_stop(old, WD_TAB_ALIGNMENT.CLEAR)
        fp.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        para_border(fp, "top", RULE, size=4, space=6)
        lr = fp.add_run(valves.header_label or "")
        lr.font.size, lr.font.color.rgb = Pt(8.5), rgb(MUTED)
        pr = fp.add_run("\tPage ")
        pr.font.size, pr.font.color.rgb = Pt(8.5), rgb(MUTED)
        field(fp, "PAGE", 8.5, MUTED)
        mid = fp.add_run(" of ")
        mid.font.size, mid.font.color.rgb = Pt(8.5), rgb(MUTED)
        field(fp, "NUMPAGES", 8.5, MUTED)

    # ---- body ----------------------------------------------------------------------------------
    blocks = parse_blocks(content)
    counters = {}

    def spacer(pts=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("")
        r.font.size = Pt(pts)
        p.paragraph_format.line_spacing = Pt(pts + 2)

    def width_weights(rows, ncols):
        w = []
        for c in range(ncols):
            longest = max((len(plain(r[c])) if c < len(r) else 0) for r in rows)
            w.append(clamp(longest, 7, 42))
        total = sum(w)
        return [x / total for x in w]

    first_h1_skipped = False
    for b in blocks:
        kind = b[0]
        if kind != "li":
            counters = {}
        if kind == "heading":
            level, text = b[1], b[2]
            if level == 1 and title and not first_h1_skipped and plain(text).strip().lower() == plain(title).strip().lower():
                first_h1_skipped = True
                continue
            h = doc.add_heading(level=min(level, 3))
            add_runs(h, text)
            if level == 1:
                para_border(h, "bottom", RULE, size=6, space=3)
        elif kind == "para":
            add_runs(doc.add_paragraph(), b[1])
        elif kind == "li":
            _, level, ordered, text = b
            if ordered:
                counters = {k: v for k, v in counters.items() if k <= level}
                counters[level] = counters.get(level, 0) + 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 + 0.8 * level)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9 + 0.8 * level))
                nr = p.add_run(f"{counters[level]}.\t")
                nr.font.bold = True
                nr.font.color.rgb = rgb(accent)
            else:
                counters = {}
                p = doc.add_paragraph(style=["List Bullet", "List Bullet 2", "List Bullet 3"][level])
            add_runs(p, text)
        elif kind == "rule":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            para_border(p, "bottom", RULE, size=6, space=1)
        elif kind == "code":
            for line in (b[1].split("\n") or [""]):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.3)
                shade(p._p.get_or_add_pPr(), "F1F3F5")
                r = p.add_run(line or " ")
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt(9.5)
            spacer()
        elif kind == "quote":
            lead = plain(b[1]).lower()
            warn = lead.startswith(("warning", "important", "caution", "risk"))
            fill, edge = (WARM_FILL, WARM_EDGE) if warn else (TINT, accent)
            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            cell = t.rows[0].cells[0]
            cell.width = text_width
            cell_shade(cell, fill)
            cell_borders(cell, left=(36, edge))
            cell_margins(t, 110, 110, 220, 180)
            row_flags(t.rows[0])
            parts = [x for x in b[1].split("\n") if x.strip()] or [""]
            for k, part in enumerate(parts):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2 if k < len(parts) - 1 else 0)
                add_runs(p, part.lstrip("-*• ").strip() if part.startswith(("- ", "* ")) else part, size=11)
            spacer(6)
        elif kind == "table":
            rows, aligns = b[1], b[2]
            ncols = max(len(r) for r in rows)
            rows = [r + [""] * (ncols - len(r)) for r in rows]
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            table_borders(t)
            cell_margins(t)
            weights = width_weights(rows, ncols)
            for ri, r in enumerate(rows):
                row = t.rows[ri]
                row_flags(row, header=(ri == 0))
                for ci, val in enumerate(r):
                    cell = row.cells[ci]
                    cell.width = int(text_width * weights[ci])
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.1
                    a = aligns[ci] if ci < len(aligns) else "l"
                    numeric = ri > 0 and NUMERIC.match(val or "") and any(ch.isdigit() for ch in val)
                    if a == "c":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif a == "r" or numeric:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if ri == 0:
                        cell_shade(cell, accent)
                        add_runs(p, val, size=10.5, color="FFFFFF", bold=True)
                    else:
                        if ri % 2 == 0:
                            cell_shade(cell, ZEBRA)
                        add_runs(p, val, size=10.5)
                        if ci == 0 and len(rows) > 2 and plain(val).lower().startswith("total"):
                            for rr in p.runs:
                                rr.font.bold = True
            spacer(8)

    cp = doc.core_properties
    cp.title, cp.author, cp.subject = plain(title) if title else "", author or "", "Created with the family AI hub"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _spacing(OxmlElement, qn, val):
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(val))
    return el


# =====================================================================================================
# Excel
# =====================================================================================================
SAFE_FUNCS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ROUNDUP", "ROUNDDOWN", "IF", "ABS", "MEDIAN", "PRODUCT", "SUMIF", "COUNTIF", "AVERAGEIF", "IFERROR", "AND", "OR", "NOT", "TODAY", "LEN"}
FORMULA_OK = re.compile(r"^=[A-Za-z0-9_+\-*/^().,:$ %<>=&\"!]*$")
CURRENCY = re.compile(r"^(?P<pre>[$€£¥]|CHF|USD|EUR)?\s?(?P<neg>-)?\s?(?P<pre2>[$€£¥])?(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)\s?(?P<post>CHF|USD|EUR|€)?$", re.I)


def safe_formula(s: str) -> bool:
    if not FORMULA_OK.match(s) or "!" in s:
        return False
    return all(f.upper() in SAFE_FUNCS for f in re.findall(r"([A-Za-z][A-Za-z0-9.]*)\(", s))


def coerce(raw: str):
    """Returns (value, number_format or None, is_number)."""
    s = (raw or "").strip()
    if not s:
        return None, None, False
    s = plain(s) if not s.startswith("=") else s
    if s.startswith("="):
        return (s, None, True) if safe_formula(s) else (s, None, "text")  # unsafe formula: kept as plain text
    m = re.match(r"^(-?\d+(?:\.\d+)?)\s?%$", s)
    if m:
        dec = len(m.group(1).split(".")[1]) if "." in m.group(1) else 0
        return float(m.group(1)) / 100, "0." + "0" * dec + "%" if dec else "0%", True
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        try:
            return dt.datetime.strptime(s, "%Y-%m-%d"), "yyyy-mm-dd", False
        except ValueError:
            pass
    if re.match(r"^0\d", s) or len(re.sub(r"\D", "", s)) > 15:
        return s, None, False
    m = CURRENCY.match(s)
    if m and (m.group("pre") or m.group("pre2") or m.group("post") or True):
        num = float(m.group("num").replace(",", ""))
        if m.group("neg"):
            num = -num
        sym = (m.group("pre") or m.group("pre2") or m.group("post") or "").upper()
        dec = "." in m.group("num")
        base = "#,##0.00" if dec else "#,##0"
        if sym in ("$", "€", "£", "¥"):
            fmt = f'"{sym}"{base}'
        elif sym:
            fmt = f'"{sym} "{base}'
        else:
            fmt = base if ("," in m.group("num") or dec) else "General"
        if not dec and not sym and "," not in m.group("num"):
            return (int(num) if num == int(num) else num), None, True
        return num, fmt, True
    return s, None, False


def build_xlsx(title, content, author, valves) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    accent, hi = valves.accent_color, valves.highlight_color
    blocks = parse_blocks(content)
    level = 2 if any(b[0] == "heading" and b[1] == 2 for b in blocks) else 1
    wb_title, sheets, cur = None, [], None
    for idx, b in enumerate(blocks):
        if b[0] == "heading" and b[1] == level:
            cur = {"name": plain(b[2]), "items": []}
            sheets.append(cur)
        elif b[0] == "heading" and b[1] < level and wb_title is None and not sheets:
            wb_title = plain(b[2])
        else:
            if cur is None:
                cur = {"name": "Sheet1", "items": []}
                sheets.append(cur)
            cur["items"].append(b)
    if not sheets:
        sheets = [{"name": "Sheet1", "items": []}]
    # fallback: no pipe tables at all -> comma/tab separated lines
    if not any(i[0] == "table" for s in sheets for i in s["items"]):
        for s in sheets:
            rows = []
            for it in s["items"]:
                if it[0] == "para":
                    rows.append([c.strip() for c in re.split(r"\t|,", it[1])])
            if rows:
                s["items"] = [("table", rows, [])]

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color=RULE)
    used = set()
    palette = [accent, hi, GREEN, "E0A12B", "7D5BA6"]
    for si, sh in enumerate(sheets):
        name = re.sub(r"[\[\]\\/*?:]", " ", sh["name"]).strip()[:31] or f"Sheet{si + 1}"
        base, n = name, 2
        while name.lower() in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = palette[si % len(palette)]
        ws.sheet_view.showGridLines = False
        r = 1  # the first table always starts at row 1, so cell references like B2 mean what the model expects
        freeze_at, filter_ref, widths = None, None, {}
        items = sh["items"]
        first_t = next((k for k, it in enumerate(items) if it[0] == "table"), None)
        if first_t:  # notes written above the first table go below it instead
            items = items[first_t:] + [("spacer",)] + items[:first_t]
        first_table = True
        for it in items:
            kind = it[0]
            if kind == "spacer":
                r += 1
            elif kind == "heading":
                c = ws.cell(row=r, column=1, value=plain(it[2]))
                c.font = Font(name="Calibri", size=12, bold=True, color=accent)
                r += 1
            elif kind == "para":
                c = ws.cell(row=r, column=1, value=plain(it[1]))
                c.font = Font(name="Calibri", size=10.5, italic=True, color=MUTED)
                r += 1
            elif kind == "li":
                c = ws.cell(row=r, column=1, value="•  " + plain(it[3]))
                c.font = Font(name="Calibri", size=10.5, color=INK)
                r += 1
            elif kind == "quote":
                c = ws.cell(row=r, column=1, value=plain(it[1]))
                c.font = Font(name="Calibri", size=10.5, bold=True, color=accent)
                c.fill = PatternFill("solid", fgColor=TINT)
                r += 1
            elif kind == "table":
                rows = it[1]
                ncols = max(len(x) for x in rows)
                rows = [x + [""] * (ncols - len(x)) for x in rows]
                top = r
                coerced = [[coerce(v) for v in rw] for rw in rows]
                num_cols, col_fmt = set(), {}
                for ci in range(ncols):
                    cells = [coerced[ri][ci] for ri in range(1, len(rows))]
                    if sum(1 for v, f, n in cells if n is True or isinstance(v, dt.datetime)) * 2 > max(1, len(cells)):
                        num_cols.add(ci)
                    fmts = [f for v, f, n in cells if f and f != "yyyy-mm-dd" and "%" not in f]
                    col_fmt[ci] = next((f for f in fmts if ".00" in f), fmts[0] if fmts else None)
                for ri, rowvals in enumerate(rows):
                    for ci, raw in enumerate(rowvals):
                        c = ws.cell(row=r, column=ci + 1)
                        if ri == 0:
                            c.value = plain(raw)
                            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                            c.fill = PatternFill("solid", fgColor=accent)
                            c.alignment = Alignment(horizontal="right" if ci in num_cols else "left", vertical="center", wrap_text=True)
                            c.border = Border(bottom=Side(style="medium", color=hi))
                        else:
                            val, fmt, isnum = coerced[ri][ci]
                            if fmt is None and isinstance(val, str) and val.startswith("=") and isnum is True:
                                fmt = col_fmt.get(ci)  # a formula cell takes its column's number format
                            c.value = val
                            if isnum == "text":
                                c.data_type = "s"
                                isnum = False
                            if fmt:
                                c.number_format = fmt
                            c.font = Font(name="Calibri", size=11, color=INK)
                            c.alignment = Alignment(horizontal="right" if isnum or isinstance(val, dt.datetime) else "left", vertical="top", wrap_text=True)
                            c.border = Border(bottom=thin)
                            if ri % 2 == 0:
                                c.fill = PatternFill("solid", fgColor=ZEBRA)
                            if ci == 0 and len(rows) > 2 and str(raw).strip().lower().startswith("total"):
                                for cc in ws[r]:
                                    cc.font = Font(name="Calibri", size=11, bold=True, color=INK)
                                    cc.border = Border(top=Side(style="medium", color=accent), bottom=thin)
                                    cc.fill = PatternFill("solid", fgColor=TINT)
                        shown = str(c.value) if c.value is not None else ""
                        widths[ci + 1] = max(widths.get(ci + 1, 0), min(len(shown.split("\n")[0]) + 2, 48))
                    r += 1
                if first_table:
                    freeze_at = f"A{top + 1}"
                    filter_ref = f"A{top}:{get_column_letter(ncols)}{r - 1}"
                    first_table = False
                r += 1
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(w, 9)
        if freeze_at:
            ws.freeze_panes = freeze_at
        if filter_ref and len(widths) > 0:
            ws.auto_filter.ref = filter_ref
        ws.page_setup.orientation = "landscape" if len(widths) > 5 else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.right.text = valves.header_label or ""
    wb.properties.title = wb_title or (plain(title) if title else "")  # the title lives in the file properties, not in a cell
    wb.properties.creator = author or "family-ai-hub"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =====================================================================================================
# PowerPoint
# =====================================================================================================
def build_pptx(title, subtitle, content, author, valves) -> bytes:
    import math

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree

    accent, hi = valves.accent_color, valves.highlight_color
    C = lambda h: RGBColor.from_string(h.upper())
    NAVY = "17384F"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    deck_title = plain(title) if title else "Presentation"

    def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = C(fill)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def add_text(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        return tb, tf

    def runs(par, text, size, color, font="Calibri", bold=False, italic=False):
        for t, f in inline_runs(text):
            r = par.add_run()
            r.text = t
            r.font.size, r.font.name = Pt(size), font
            r.font.bold = bool(bold or f.get("b"))
            r.font.italic = bool(italic or f.get("i"))
            r.font.color.rgb = C(color)
            if f.get("url"):
                r.hyperlink.address = f["url"]
            if f.get("code"):
                r.font.name = "Consolas"

    def bullet(par, level, ordered, size):
        ppr = par._p.get_or_add_pPr()
        indent = Inches(0.38)
        ppr.set("marL", str(int(Inches(0.38) + level * Inches(0.45))))
        ppr.set("indent", str(-int(indent)))
        clr = etree.SubElement(ppr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", hi if level == 0 else accent)
        if ordered:
            etree.SubElement(ppr, qn("a:buFont")).set("typeface", "+mj-lt")
            etree.SubElement(ppr, qn("a:buAutoNum")).set("type", "arabicPeriod")
        else:
            etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(ppr, qn("a:buChar")).set("char", "•" if level == 0 else "–")

    def notes(slide, text):
        if text.strip():
            slide.notes_slide.notes_text_frame.text = text.strip()

    layout_title_only = prs.slide_layouts[5]
    layout_blank = prs.slide_layouts[6]
    count = {"n": 0}

    # ---- title slide with a ridge line ---------------------------------------------------------
    def title_slide():
        s = prs.slides.add_slide(layout_blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = C(NAVY)

        def ridge(points, color):
            fb = s.shapes.build_freeform(int(points[0][0] * W), int(points[0][1] * H))
            fb.add_line_segments([(int(x * W), int(y * H)) for x, y in points[1:]], close=True)
            shp = fb.convert_to_shape()
            shp.fill.solid()
            shp.fill.fore_color.rgb = C(color)
            shp.line.fill.background()
            shp.shadow.inherit = False
            return shp

        ridge([(0, 1), (0, .78), (.07, .70), (.12, .74), (.2, .60), (.27, .68), (.34, .64), (.44, .76), (.54, .66), (.62, .72), (.7, .58), (.78, .68), (.86, .62), (1, .74), (1, 1)], "245B7E")
        ridge([(0, 1), (0, .90), (.1, .85), (.2, .88), (.33, .82), (.46, .89), (.6, .84), (.75, .90), (.88, .85), (1, .89), (1, 1)], GREEN)
        # small Swiss flag on the highest peak
        fx, fy = int(.70 * W), int(.58 * H)
        pole = rect(s, fx - Inches(0.01), fy - Inches(0.46), Inches(0.02), Inches(0.46), "E9EEF3")
        flag = rect(s, fx + Inches(0.02), fy - Inches(0.46), Inches(0.26), Inches(0.26), "D52B1E")
        rect(s, fx + Inches(0.13) - Inches(0.015), fy - Inches(0.46) + Inches(0.05), Inches(0.03), Inches(0.16), "FFFFFF")
        rect(s, fx + Inches(0.075), fy - Inches(0.46) + Inches(0.115), Inches(0.12), Inches(0.03), "FFFFFF")
        rect(s, Inches(0.9), Inches(1.7), Inches(0.9), Inches(0.09), hi)
        _, tf = add_text(s, Inches(0.85), Inches(0.9), Inches(9), Inches(0.5))
        runs(tf.paragraphs[0], (valves.header_label or "").upper(), 13, "9FC3DC", bold=True)
        _, tf = add_text(s, Inches(0.85), Inches(2.0), Inches(11.4), Inches(2.0), MSO_ANCHOR.TOP)
        runs(tf.paragraphs[0], deck_title, 44, "FFFFFF", "Georgia", bold=True)
        if subtitle:
            p = tf.add_paragraph()
            p.space_before = Pt(10)
            runs(p, plain(subtitle), 22, "CFE0EC")
        meta = " · ".join(x for x in (author, f"{dt.date.today().day} {dt.date.today():%B %Y}") if x)
        _, tf = add_text(s, Inches(0.85), Inches(3.75), Inches(8), Inches(0.4))
        runs(tf.paragraphs[0], meta, 14, "9FC3DC")
        return s

    def chrome(slide, title_text):
        count["n"] += 1
        rect(slide, 0, 0, Inches(0.22), H, accent)
        rect(slide, 0, Inches(0.55), Inches(0.22), Inches(0.9), hi)
        t = slide.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.8), Inches(0.5), Inches(11.8), Inches(1.0)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        par = tf.paragraphs[0]
        par.alignment = PP_ALIGN.LEFT
        size = 34 if len(title_text) <= 40 else 28 if len(title_text) <= 70 else 24
        runs(par, title_text, size, INK, "Georgia", bold=True)
        ln = rect(slide, Inches(0.85), Inches(1.55), Inches(11.7), Pt(1.2), RULE)
        _, ftf = add_text(slide, Inches(0.85), Inches(6.98), Inches(9), Inches(0.35))
        runs(ftf.paragraphs[0], deck_title, 10.5, MUTED)
        _, ntf = add_text(slide, Inches(11.6), Inches(6.98), Inches(1.0), Inches(0.35))
        ntf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        runs(ntf.paragraphs[0], str(count["n"] + 1), 10.5, MUTED)

    def new_slide(title_text):
        s = prs.slides.add_slide(layout_title_only)
        chrome(s, title_text)
        return s

    def section_slide(text):
        s = prs.slides.add_slide(layout_blank)
        count["n"] += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = C(accent)
        rect(s, Inches(0.9), Inches(3.0), Inches(0.9), Inches(0.09), "FFFFFF")
        _, tf = add_text(s, Inches(0.85), Inches(3.3), Inches(11.4), Inches(2.0))
        runs(tf.paragraphs[0], text, 40, "FFFFFF", "Georgia", bold=True)
        return s

    def flow_slide(title_text, items, note):
        """items: ('para', text) ('li', level, ordered, text) ('quote', text)."""
        # budget lines per slide: ~66 characters per line at 24pt across the text area
        def lines_for(it):
            text = plain(it[-1])
            size = 28 if it[0] != "li" or it[1] == 0 else 24
            per = 57 if size == 28 else 66
            return max(1, math.ceil(len(text) / per)) + (0.45 if it[0] == "li" else 0.6)
        chunks, cur, used = [], [], 0.0
        for it in items:
            need = lines_for(it)
            if cur and used + need > 8.0:
                chunks.append(cur)
                cur, used = [], 0.0
            cur.append(it)
            used += need
        if cur:
            chunks.append(cur)
        for k, chunk in enumerate(chunks or [[]]):
            s = new_slide(title_text + (" (cont.)" if k else ""))
            _, tf = add_text(s, Inches(0.85), Inches(1.8), Inches(11.7), Inches(4.95))
            first = True
            for it in chunk:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.space_after = Pt(9)
                p.line_spacing = 1.05
                if it[0] == "li":
                    size = 28 if it[1] == 0 else 24
                    runs(p, it[3], size, INK if it[1] == 0 else "3B4856")
                    bullet(p, it[1], it[2], size)
                elif it[0] == "quote":
                    runs(p, it[1], 28, accent, italic=True)
                else:
                    runs(p, it[1], 28, INK)
            if k == 0:
                notes(s, note)

    def key_message_slide(title_text, text, note):
        s = new_slide(title_text)
        rect(s, Inches(0.85), Inches(2.1), Inches(11.7), Inches(3.6), TINT)
        rect(s, Inches(0.85), Inches(2.1), Inches(0.12), Inches(3.6), hi)
        _, tf = add_text(s, Inches(1.35), Inches(2.3), Inches(10.8), Inches(3.2), MSO_ANCHOR.MIDDLE)
        runs(tf.paragraphs[0], text, 30, INK, "Georgia", italic=False)
        notes(s, note)

    def table_slides(title_text, rows, intro, note):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        header, body = rows[0], rows[1:]
        per = 8
        pages = [body[i : i + per] for i in range(0, len(body), per)] or [[]]
        weights = [clamp(max(len(plain(r[c])) for r in rows), 6, 38) for c in range(ncols)]
        num_cols = {c for c in range(ncols) if sum(1 for r in body if NUMERIC.match(r[c] or "") and any(ch.isdigit() for ch in r[c])) * 2 > max(1, len(body))}
        tot = sum(weights)
        for k, page in enumerate(pages):
            s = new_slide(title_text + (" (cont.)" if k else ""))
            y = Inches(1.8)
            if intro and k == 0:
                _, tf = add_text(s, Inches(0.85), y, Inches(11.7), Inches(0.7))
                runs(tf.paragraphs[0], intro, 20, "3B4856")
                y += Inches(0.8)
            n = len(page) + 1
            size = 20 if n <= 6 else 18 if n <= 8 else 15
            rh = Inches(0.62 if n <= 6 else 0.52 if n <= 8 else 0.46)
            gt = s.shapes.add_table(n, ncols, Inches(0.85), y, Inches(11.7), rh * n)
            tbl = gt.table
            tblPr = gt._element.graphic.graphicData.tbl.tblPr
            for attr in ("bandRow", "firstRow"):
                tblPr.set(attr, "0")
            style = tblPr.find(qn("a:tableStyleId"))
            if style is not None:
                style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
            for c in range(ncols):
                tbl.columns[c].width = int(Inches(11.7) * weights[c] / tot)
            for ri, rowvals in enumerate([header] + page):
                tbl.rows[ri].height = rh
                for ci, val in enumerate(rowvals):
                    cell = tbl.cell(ri, ci)
                    cell.margin_left = cell.margin_right = Inches(0.12)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = C(accent if ri == 0 else (ZEBRA if ri % 2 == 0 else "FFFFFF"))
                    par = cell.text_frame.paragraphs[0]
                    par.alignment = PP_ALIGN.RIGHT if ci in num_cols else PP_ALIGN.LEFT
                    runs(par, val, size, "FFFFFF" if ri == 0 else INK, bold=(ri == 0))
            if k == 0:
                notes(s, note)

    # ---- assemble slides -----------------------------------------------------------------------
    title_slide()
    blocks = parse_blocks(content)
    slides, cur = [], None
    for b in blocks:
        if b[0] == "heading" and b[1] == 1:
            if not slides and plain(b[2]).strip().lower() == deck_title.strip().lower():
                continue
            slides.append({"kind": "section", "title": plain(b[2]), "items": []})
            cur = None
        elif b[0] == "heading" and b[1] >= 2:
            cur = {"kind": "slide", "title": plain(b[2]), "items": []}
            slides.append(cur)
        else:
            if cur is None:
                cur = {"kind": "slide", "title": deck_title if not slides else "Overview", "items": []}
                slides.append(cur)
            cur["items"].append(b)

    for sl in slides:
        if sl["kind"] == "section":
            section_slide(sl["title"])
            continue
        note_lines, items, in_notes = [], [], False
        for it in sl["items"]:
            if it[0] == "para" and re.match(r"^(speaker )?notes?:", it[1], re.I):
                in_notes = True
                note_lines.append(re.sub(r"^(speaker )?notes?:\s*", "", it[1], flags=re.I))
            elif in_notes and it[0] in ("para", "li"):
                note_lines.append(it[-1])
            else:
                items.append(it)
        note = "\n".join(note_lines)
        tables = [i for i in items if i[0] == "table"]
        flow = [i for i in items if i[0] in ("para", "li", "quote")]
        if tables:
            intro = " ".join(plain(i[1]) for i in flow if i[0] == "para")[:200]
            table_slides(sl["title"], tables[0][1], intro, note)
        elif flow and all(i[0] == "quote" for i in flow) and len(flow) == 1 and len(plain(flow[0][1])) < 260:
            key_message_slide(sl["title"], plain(flow[0][1]), note)
        elif flow or True:
            code = [i[1] for i in items if i[0] == "code"]
            flow = flow + [("para", c) for c in code]
            flow_slide(sl["title"], flow, note)

    prs.core_properties.title = deck_title
    prs.core_properties.author = author or "family-ai-hub"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# =====================================================================================================
# The Open WebUI tool
# =====================================================================================================
class Tools:
    class Valves(BaseModel):
        """Admin settings."""

        header_label: str = Field(default="Buzatu · Miller family", description="Small label printed in headers, footers and on title slides.")
        accent_color: str = Field(default="2A6A98", description="Main colour as 6 hex digits (lake blue).")
        highlight_color: str = Field(default="D52B1E", description="Accent for rules and highlights as 6 hex digits (alpine red).")
        paper_size: Literal["A4", "Letter"] = Field(default="A4", description="Page size for Word documents.")
        max_input_chars: int = Field(default=150000, description="Longest text accepted for one file.")

    def __init__(self):
        self.valves = self.Valves()
        self.citation = False  # the result is a download link, not a source

    # ----- helpers --------------------------------------------------------------------------------
    def _clean(self, text: str) -> str:
        text = text or ""
        if len(text) > self.valves.max_input_chars:
            raise ValueError(f"The text is too long ({len(text)} characters; the limit is {self.valves.max_input_chars}). Split it into several files.")
        return text

    async def _save(self, data: bytes, filename: str, mime: str, user: dict, emitter) -> str:
        from open_webui.models.files import FileForm, Files
        from open_webui.storage.provider import Storage

        user = user or {}
        file_id = str(uuid.uuid4())
        stored_name = f"{file_id}_{filename}"
        tags = {"OpenWebUI-User-Id": str(user.get("id", "")), "OpenWebUI-File-Id": file_id}
        _, path = await asyncio.to_thread(Storage.upload_file, io.BytesIO(data), stored_name, tags)
        res = Files.insert_new_file(
            user.get("id"),
            FileForm(id=file_id, filename=filename, path=path, data={}, meta={"name": filename, "content_type": mime, "size": len(data), "data": {"created_by": "office_export"}}),
        )
        if inspect.isawaitable(res):
            res = await res
        if not res:
            raise RuntimeError("Open WebUI could not register the file.")
        link = f"/api/v1/files/{file_id}/content?attachment=true"
        if emitter:
            await emitter({"type": "status", "data": {"description": f"Created {filename}", "done": True}})
        return (
            f"Done. The file '{filename}' has been created ({len(data) // 1024 or 1} KB). "
            f"Tell the user it is ready and show this download link exactly as written, without changing the address: "
            f"[Download {filename}]({link}). Do not repeat the contents of the file."
        )

    # ----- Word -----------------------------------------------------------------------------------
    async def create_word_document(self, title: str, content: str, subtitle: str = "", __user__: Optional[dict] = None, __event_emitter__=None) -> str:
        """
        Create a polished Word (.docx) document and return a download link.
        Use it when the user asks for a document, report, briefing, letter, memo, meeting notes or handout as a Word file.

        Write `content` in Markdown. The tool turns it into a designed document, so do not add your own page furniture.
        Supported: # ## ### headings, normal paragraphs, **bold**, *italic*, `code`, [links](https://...),
        - bullet lists (indent two spaces for sub-points), 1. numbered lists, | tables | with a header row |,
        > quotes (shown as a shaded call-out box; start the quote with "Warning:" or "Important:" for an orange box),
        --- horizontal rules and ``` code blocks.
        Do not put the title in `content` (it is printed automatically from `title`). Start with the first section heading.

        Tips for an attractive result:
        - Open with a short "Summary" or "Key takeaways" section of 3 to 5 bullets, or a "> **Key takeaway:** ..." call-out.
        - Use ## for sections and ### for sub-sections. Keep headings short.
        - Keep paragraphs to 2 to 4 sentences. Prefer bullets for lists of facts.
        - Use a table for anything with several attributes per item (comparisons, schedules, budgets).
        - Put the sources at the end under "## Sources" as a bullet list of links.

        :param title: The document title, for example "Briefing: The Swiss Cloud Debate".
        :param content: The body of the document in Markdown, without the title.
        :param subtitle: Optional one-line subtitle shown under the title.
        :return: A message containing the download link to show to the user.
        """
        try:
            user = __user__ or {}
            data = await asyncio.to_thread(build_docx, title, subtitle, self._clean(content), user.get("name", ""), self.valves)
            return await self._save(data, f"{slug(title)}.docx", DOCX_MIME, user, __event_emitter__)
        except Exception as e:  # reported to the model so it can fix its input
            return f"Could not create the Word document: {e}"

    # ----- Excel ----------------------------------------------------------------------------------
    async def create_excel_workbook(self, title: str, content: str, __user__: Optional[dict] = None, __event_emitter__=None) -> str:
        """
        Create a formatted Excel (.xlsx) workbook and return a download link.
        Use it when the user asks for a spreadsheet, table, budget, tracker, schedule or list as an Excel file.

        Write `content` in Markdown. Start each worksheet with a "## Sheet name" heading, followed by a Markdown table
        (| Column A | Column B | on the first row, then a |---|---| line, then the data rows).
        Numbers, percentages (45%), currency ($1,200.50, CHF 80, €15) and dates (2026-10-21) become real Excel values with
        proper formatting. A row whose first cell starts with "Total" is styled as a totals row.
        Formulas are allowed (SUM, AVERAGE, MIN, MAX, COUNT, ROUND, IF and simple arithmetic). Cell references count from the first
        table on the sheet: its header row is row 1, the first data row is row 2, and the columns are A, B, C, ... Use formulas only
        with one table per sheet. Plain text lines become notes placed below the table. Do not write numbers as words or approximations.

        Tips: one table per sheet, a short clear header row, one fact per cell, consistent units in a column.

        :param title: The workbook title, for example "Household Budget 2026".
        :param content: The sheets as Markdown: "## Sheet name" headings, each followed by a table.
        :return: A message containing the download link to show to the user.
        """
        try:
            user = __user__ or {}
            data = await asyncio.to_thread(build_xlsx, title, self._clean(content), user.get("name", ""), self.valves)
            return await self._save(data, f"{slug(title, 'spreadsheet')}.xlsx", XLSX_MIME, user, __event_emitter__)
        except Exception as e:
            return f"Could not create the Excel workbook: {e}"

    # ----- PowerPoint -----------------------------------------------------------------------------
    async def create_powerpoint(self, title: str, content: str, subtitle: str = "", __user__: Optional[dict] = None, __event_emitter__=None) -> str:
        """
        Create a designed PowerPoint (.pptx) presentation and return a download link.
        Use it when the user asks for slides, a deck or a presentation.

        Write `content` in Markdown. A title slide is created from `title` and `subtitle`; do not repeat them.
        Each "## Slide title" starts a new slide. A "# Section name" heading makes a full-colour divider slide.
        Under a slide title use short bullets ("- point", two spaces to indent a sub-point), or one table, or a single
        "> key message" line, which becomes a large call-out slide. A paragraph starting with "Notes:" becomes speaker notes
        (not shown on the slide). Slides with too much text are split automatically.

        Tips for attractive slides: 3 to 5 bullets per slide, each under 12 words; one idea per slide; start with an agenda or
        summary slide and end with a "Next steps" or "Key takeaways" slide; put detail in the speaker notes, not on the slide.

        :param title: The presentation title.
        :param content: The slides as Markdown.
        :param subtitle: Optional subtitle shown on the title slide.
        :return: A message containing the download link to show to the user.
        """
        try:
            user = __user__ or {}
            data = await asyncio.to_thread(build_pptx, title, subtitle, self._clean(content), user.get("name", ""), self.valves)
            return await self._save(data, f"{slug(title, 'presentation')}.pptx", PPTX_MIME, user, __event_emitter__)
        except Exception as e:
            return f"Could not create the presentation: {e}"
